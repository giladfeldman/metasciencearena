"""citeclaim Stage E: frontier panel over the MetaMisCitations pilot. GOLD-BUILDING code.

Inputs and outputs are real article text (citing sentences, abstracts, quotes), so
all of it lives in article custody, never in this repository:
    $ARTICLE_REPOSITORY/ground_truth/citeclaim-2026-10/stage_e/
        jobs/*.json          written by MetaMisCitations (src/citeclaim_stage, `prepare`)
        results/             the free model's citeclaim results (NOT read by `panel`)
        panel/<seat>.jsonl   one record per call, appended, resumable
        panel/consensus.json per-item verdicts of the three seats + the free model

Stages:
  panel --seat opus|sonnet|sol
        Every citation citeclaim would send to a model (the library's eligibility,
        recomputed here from the job inputs alone, so the panel is BLIND to the
        free model's verdicts) is judged by one frontier seat. Calls are grouped
        exactly as the library groups them (per source, document order, chunks of
        at most 6), and the prompt is prompts/v1.md rendered with str.format, which
        is byte-identical to citeclaim's claim-pair-v1 (checked against the
        library's renderedPromptSha256 by `consensus`).
          opus, sonnet: `claude --print --model <alias>` on Claude Max, hardened by
                        framework.hermetic (no tools, scratch cwd, no settings).
          sol:          `codex exec -m gpt-6.1-sol`, hardened by framework.hermetic;
                        nonce-probed 2026-10-18 (runs/stage_e/codex_nonce_probe_2026-10-18.json).
        Answers are guarded exactly as the arena scorer guards them (a consistent
        or contradicts verdict whose quote is not in the abstract is no verdict).
  consensus
        Joins the three seats and the free model per item. Writes panel/consensus.json.

    python arenas/citation-claim-fidelity-v1/tools/stage_e_panel.py panel --seat opus
    python arenas/citation-claim-fidelity-v1/tools/stage_e_panel.py consensus
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ARENA = Path(__file__).resolve().parents[1]
REPO = ARENA.parents[1]
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(ARENA))
from framework import hermetic  # noqa: E402
import scorer  # noqa: E402
from generator import render_sentences  # noqa: E402

#: Custody folder. ARTICLE_REPOSITORY unset -> a path that does not exist, so a
#: stage fails on its first read instead of reading a wrong folder.
STAGE_E = Path(os.environ.get("ARTICLE_REPOSITORY") or "<ARTICLE_REPOSITORY unset>") / "ground_truth" / "citeclaim-2026-10" / "stage_e"
PROMPT = (ARENA / "prompts" / "v1.md").read_text(encoding="utf-8").replace("\r\n", "\n")
MIN_ABSTRACT_CHARS = 200  # citeclaim eligibility.MIN_ABSTRACT_CHARS
MAX_GROUP = 6

SEATS = {
    "opus": {"argv": ["claude", "--print", "--model", "opus", "--output-format", "json"], "family": "anthropic"},
    "sonnet": {"argv": ["claude", "--print", "--model", "sonnet", "--output-format", "json"], "family": "anthropic"},
    "sol": {"argv": ["codex", "exec", "-m", "gpt-6.1-sol", "-c", 'model_reasoning_effort="high"'], "family": "openai"},
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _sha(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def normalize_doi(d):
    if not d:
        return None
    d = re.sub(r"^https?://(dx\.)?doi\.org/", "", d.strip(), flags=re.I)
    d = re.sub(r"^doi:\s*", "", d, flags=re.I).lower()
    return d if d.startswith("10.") else None


def eligible(c: dict, abstracts: dict) -> bool:
    """citeclaim ineligibilityReason(c) is None."""
    if c.get("linkStatus") != "matched" or not c.get("referenceKey"):
        return False
    if c.get("sourceCount") != 1:
        return False
    doi = normalize_doi(c.get("doi"))
    if not doi or not (c.get("sentence") or "").strip():
        return False
    a = abstracts.get(doi)
    if not a or not a.get("abstract") or len(a["abstract"].strip()) < MIN_ABSTRACT_CHARS:
        return False
    return not (a.get("doi") and normalize_doi(a["doi"]) != doi)


def groups() -> list[dict]:
    """Every call the library makes on the jobs, in the library's order."""
    out = []
    for jf in sorted((STAGE_E / "jobs").glob("*.json")):
        job = json.loads(jf.read_text(encoding="utf-8"))
        abstracts = {a["doi"]: a for a in job["abstracts"]}
        by_src: dict[str, list[dict]] = {}
        for c in job["citations"]:
            if eligible(c, abstracts):
                by_src.setdefault(json.dumps([c["referenceKey"], normalize_doi(c["doi"])]), []).append(c)
        for cites in by_src.values():
            for i in range(0, len(cites), MAX_GROUP):
                chunk = cites[i:i + MAX_GROUP]
                a = abstracts[normalize_doi(chunk[0]["doi"])]
                prompt = PROMPT.format(title=a.get("title") or "(unknown)",
                                       year=a.get("year") if a.get("year") not in (None, "") else "n.d.",
                                       abstract=a["abstract"],
                                       sentences=render_sentences([c["sentence"] for c in chunk]))
                out.append({"job": jf.stem, "doi": a["doi"], "citation_ids": [c["id"] for c in chunk],
                            "abstract": a["abstract"], "prompt": prompt, "rendered_sha256": _sha(prompt)})
    return out


def extract_json_object(text: str):
    """citeclaim extractJsonObject: fenced block if any, then first '{' .. last '}'."""
    s = (text or "").strip()
    m = re.search(r"```(?:json)?\s*([\s\S]*?)```", s, re.I)
    if m:
        s = m.group(1).strip()
    i, j = s.find("{"), s.rfind("}")
    if i < 0 or j <= i:
        raise ValueError("no JSON object in answer")
    return json.loads(s[i:j + 1])


def call_seat(seat: str, prompt: str) -> dict:
    argv, kw = hermetic.spawn_kwargs(SEATS[seat]["argv"])
    t0 = time.time()
    cp = subprocess.run(argv, input=prompt, capture_output=True, text=True, encoding="utf-8",
                        errors="replace", timeout=1200, **kw)
    rec = {"argv_hardened": argv, "exit": cp.returncode, "seconds": round(time.time() - t0, 1),
           "stderr_tail": cp.stderr[-600:]}
    if seat in ("opus", "sonnet"):
        try:
            d = json.loads(cp.stdout)
            rec["text"] = d.get("result") or ""
            rec["is_error"] = bool(d.get("is_error"))
            rec["resolved_models"] = list((d.get("modelUsage") or {}).keys())
            rec["usage"] = d.get("usage")
            rec["cost_usd_equivalent"] = d.get("total_cost_usd")
        except json.JSONDecodeError:
            rec["text"], rec["is_error"] = "", True
            rec["stdout_tail"] = cp.stdout[-600:]
    else:
        rec["text"] = cp.stdout
        rec["is_error"] = cp.returncode != 0 or not cp.stdout.strip()
        m = re.search(r"^model:\s*(\S+)", cp.stderr, re.M)
        rec["resolved_models"] = [m.group(1)] if m else []
        m = re.search(r"^reasoning effort:\s*(\S+)", cp.stderr, re.M)
        rec["reasoning_effort"] = m.group(1) if m else None
        m = re.search(r"tokens used\s*\n\s*([\d,]+)", cp.stderr)
        rec["tokens_used"] = int(m.group(1).replace(",", "")) if m else None
    return rec


def panel(seat: str, limit: int | None) -> None:
    out = STAGE_E / "panel" / f"{seat}.jsonl"
    out.parent.mkdir(parents=True, exist_ok=True)
    done = set()
    if out.exists():
        for line in out.read_text(encoding="utf-8").splitlines():
            r = json.loads(line)
            if not r.get("call", {}).get("is_error"):
                done.add(r["rendered_sha256"])
    gs = [g for g in groups() if g["rendered_sha256"] not in done]
    if limit:
        gs = gs[:limit]
    binary = SEATS[seat]["argv"][0]
    meta = {"seat": seat, "cli_version": hermetic.cli_version(binary),
            "hermetic_verified": hermetic.HERMETIC_PROFILES[binary].verified,
            "call_config": {"argv": SEATS[seat]["argv"], "prompt": "prompts/v1.md (claim-pair-v1), str.format",
                            "prompt_sha256": _sha(PROMPT), "grouping": f"per source, chunks of {MAX_GROUP}",
                            "temperature": "cli default (not exposed)"}}
    print(f"[panel {seat}] {len(gs)} calls to make ({len(done)} already done)")
    for k, g in enumerate(gs, 1):
        rec = None
        for attempt in range(1, 4):
            rec = call_seat(seat, g["prompt"])
            if not rec["is_error"]:
                break
            time.sleep(30 * attempt)
        try:
            parsed = extract_json_object(rec["text"])
        except (ValueError, json.JSONDecodeError):
            parsed = None
        n = len(g["citation_ids"])
        guarded = scorer.guarded_verdicts(parsed, n, g["abstract"]) if parsed is not None else [None] * n
        items = (parsed or {}).get("items") if isinstance(parsed, dict) else None
        by_id = {}
        for it in items if isinstance(items, list) else []:
            if isinstance(it, dict) and scorer.parse_item_id(it.get("id")) is not None:
                by_id[scorer.parse_item_id(it.get("id"))] = it
        row = {**meta, "at": _now(), "job": g["job"], "doi": g["doi"], "rendered_sha256": g["rendered_sha256"],
               "citation_ids": g["citation_ids"], "call": {k2: v for k2, v in rec.items()},
               "verdicts": [{"citation_id": cid, "guarded": guarded[i],
                             "raw_verdict": (by_id.get(i + 1) or {}).get("verdict"),
                             "evidence_quote": (by_id.get(i + 1) or {}).get("evidence_quote"),
                             "reason": (by_id.get(i + 1) or {}).get("reason")}
                            for i, cid in enumerate(g["citation_ids"])]}
        with out.open("a", encoding="utf-8") as f:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
        print(f"[panel {seat}] {k}/{len(gs)} {rec['seconds']} s {Counter(guarded)}", flush=True)


def _latest(seat: str) -> dict:
    """citation id -> verdict record from the seat's latest non-error call."""
    p = STAGE_E / "panel" / f"{seat}.jsonl"
    out = {}
    if p.exists():
        for line in p.read_text(encoding="utf-8").splitlines():
            r = json.loads(line)
            if r["call"].get("is_error"):
                continue
            for v in r["verdicts"]:
                out[v["citation_id"]] = {**v, "resolved_models": r["call"].get("resolved_models"),
                                         "rendered_sha256": r["rendered_sha256"]}
    return out


def _free_model() -> dict:
    """citation id -> merged free-model item (r1 then retry rounds), plus the library's rendered sha per call."""
    items, shas = {}, {}
    base = STAGE_E / "results"
    for jf in sorted((STAGE_E / "jobs").glob("*.json")):
        rounds = [base / "r1" / jf.name] + sorted(base.glob(f"retry*/{jf.name}"))
        for rp in [r for r in rounds if r.exists()]:
            res = json.loads(rp.read_text(encoding="utf-8"))["result"]
            for c in res["calls"]:
                for cid in c["citationIds"]:
                    shas.setdefault(cid, set()).add(c["renderedPromptSha256"])
            for it in res["items"]:
                prev = items.get(it["citationId"])
                if prev is None or prev.get("reason") == "llm_unavailable":
                    items[it["citationId"]] = it
    return items, shas


def consensus() -> None:
    gs = groups()
    seats = {s: _latest(s) for s in SEATS}
    free, free_shas = _free_model()
    rows, prompt_match = [], Counter()
    for g in gs:
        for cid in g["citation_ids"]:
            vs = {s: (seats[s].get(cid) or {}).get("guarded") for s in SEATS}
            have = all(cid in seats[s] for s in SEATS)
            unanimous = have and len(set(vs.values())) == 1 and None not in vs.values()
            fm = free.get(cid) or {}
            prompt_match[g["rendered_sha256"] in free_shas.get(cid, set())] += 1
            rows.append({"citation_id": cid, "job": g["job"], "doi": g["doi"], "panel": vs,
                         "panel_complete": have, "unanimous": unanimous,
                         "consensus": vs["opus"] if unanimous else None,
                         "free_model": {"status": fm.get("status"), "verdict": fm.get("verdict"),
                                        "reason": fm.get("reason")}})
    out = {"built_at": _now(), "items": rows,
           "summary": {"items": len(rows), "panel_complete": sum(r["panel_complete"] for r in rows),
                       "unanimous": sum(r["unanimous"] for r in rows),
                       "consensus_verdicts": dict(Counter(r["consensus"] for r in rows if r["unanimous"])),
                       "same_prompt_as_library": dict(prompt_match)}}
    p = STAGE_E / "panel" / "consensus.json"
    p.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(out["summary"], indent=1))


SEED = 20261018
DECISIONS = ("consistent", "contradicts", "silent", "unsure")
N_CONSISTENT_AUDIT = 20


def _largest_remainder(shares: dict, total: int) -> dict:
    s = sum(shares.values())
    if not s:
        return {k: 0 for k in shares}
    raw = {k: total * v / s for k, v in shares.items()}
    out = {k: int(v) for k, v in raw.items()}
    for k in sorted(raw, key=lambda k: (-(raw[k] - out[k]), k))[: total - sum(out.values())]:
        out[k] += 1
    return out


def select_for_owner(rows: list[dict], disc_of: dict) -> dict:
    """The pre-registered sheet (docs/design/2026-10-citeclaim-stage-e-mmc-pilot.md). Returns id -> strata."""
    import math
    import random
    rng = random.Random(SEED)
    strata: dict[str, list[str]] = {}
    no_cons = [r["citation_id"] for r in rows if not r["unanimous"]]
    unan = sorted(r["citation_id"] for r in rows if r["unanimous"])
    for cid in no_cons:
        strata.setdefault(cid, []).append("a_no_consensus")
    for cid in rng.sample(unan, math.ceil(0.10 * len(unan))) if unan else []:
        strata.setdefault(cid, []).append("b_unanimous_audit")
    for r in rows:
        if r["free_model"]["verdict"] == "contradicts":
            strata.setdefault(r["citation_id"], []).append("c_free_contradicts")
    pool = sorted(r["citation_id"] for r in rows
                  if r["free_model"]["verdict"] == "consistent" and r["citation_id"] not in strata)
    all_cons = [r["citation_id"] for r in rows if r["free_model"]["verdict"] == "consistent"]
    alloc = _largest_remainder(Counter(disc_of[c] for c in all_cons), N_CONSISTENT_AUDIT)
    for d in sorted(alloc):
        cand = [c for c in pool if disc_of[c] == d]
        for cid in rng.sample(cand, min(alloc[d], len(cand))):
            strata.setdefault(cid, []).append("c_free_consistent_sample")
    return strata


def sheet() -> None:
    """Build the owner review sheet (xlsx) and its key in custody."""
    import random
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    cons = json.loads((STAGE_E / "panel" / "consensus.json").read_text(encoding="utf-8"))
    rows = cons["items"]
    if not all(r["panel_complete"] for r in rows):
        missing = sum(not r["panel_complete"] for r in rows)
        raise SystemExit(f"panel incomplete for {missing} items; finish the panel before building the sheet")
    jobs = {jf.stem: json.loads(jf.read_text(encoding="utf-8")) for jf in (STAGE_E / "jobs").glob("*.json")}
    cit, meta, absd = {}, {}, {}
    for stem, job in jobs.items():
        for c in job["citations"]:
            cit[c["id"]] = c
        for m in job["pairs_meta"]:
            meta[m["id"]] = m
        for a in job["abstracts"]:
            absd[a["doi"]] = a
    seats = {s: _latest(s) for s in SEATS}
    strata = select_for_owner(rows, {cid: meta[cid]["discipline"] for cid in meta})
    order = sorted(strata)
    random.Random(SEED).shuffle(order)

    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    head = ["row", "decision", "note", "citing sentence", "cited as", "citing journal", "cited source", "DOI link",
            "abstract", "Opus verdict", "Opus quote / reason", "Sonnet verdict", "Sonnet quote / reason",
            "Sol verdict", "Sol quote / reason"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    key = []
    for i, cid in enumerate(order, 1):
        c, m = cit[cid], meta[cid]
        a = absd[c["doi"]]
        line = [f"R{i:03d}", None, None, c["sentence"], m["cited_ref"], m["journal"],
                f"{a.get('title') or '(no title)'} ({a.get('year') or 'n.d.'})", f"https://doi.org/{c['doi']}",
                a["abstract"]]
        for s in SEATS:
            v = seats[s][cid]
            line += [v["guarded"] or f"no verdict (said: {v['raw_verdict']})",
                     (f"“{v['evidence_quote']}” — " if v.get("evidence_quote") else "") + (v.get("reason") or "")]
        ws.append(line)
        ws.cell(row=i + 1, column=8).hyperlink = line[7]
        key.append({"row": f"R{i:03d}", "citation_id": cid, "strata": strata[cid]})
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{len(order) + 1}")
    widths = [6, 13, 20, 60, 18, 22, 40, 28, 90, 13, 45, 13, 45, 13, 45]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    info = wb.create_sheet("how to decide")
    for t in [
        "One row = one citing sentence and the abstract of the source it cites.",
        "Judge only the claim the sentence attributes to that source, using ONLY the abstract (as the models did).",
        "decision: consistent = the abstract supports the claim as stated (a faithful paraphrase counts);",
        "          contradicts = the abstract conflicts with it (other direction, number, design, measure, population, or stated more strongly/generally);",
        "          silent = the abstract neither supports nor contradicts it; unsure = cannot decide (the row is then left out of scoring).",
        "The three model verdicts are shown for convenience; they are not the answer. Rows are in random order.",
        f"Built {_now()} from panel/consensus.json; seed {SEED}; rules: ScienceArena docs/design/2026-10-citeclaim-stage-e-mmc-pilot.md.",
    ]:
        info.append([t])
    info.column_dimensions["A"].width = 150
    out = STAGE_E / "review" / "stage_e_review_sheet_r1.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    (STAGE_E / "review" / "stage_e_review_key_r1.json").write_text(
        json.dumps({"built_at": _now(), "seed": SEED, "rows": key,
                    "strata_counts": dict(Counter(s for k in key for s in k["strata"]))}, indent=1), encoding="utf-8")
    print(json.dumps({"sheet": str(out), "rows": len(key),
                      "strata": dict(Counter(s for k in key for s in k["strata"]))}, indent=1))


# ------------------------------------------------------------------ decider (amendment 1/1a)
DECIDER_ARGV = ["claude", "--print", "--model", "opus", "--effort", "xhigh", "--output-format", "json"]
DECIDER_VERDICTS = ("consistent", "contradicts", "silent", "undecidable")
DECIDER_PROMPT = """You are the deciding judge in a benchmark of citation accuracy. Three reviewers judged independently whether ONE sentence from a manuscript describes a cited source accurately, using only the cited source's abstract as evidence. Their verdicts may agree or disagree. Make the final decision yourself.

The scheme (the same one the reviewers used):
- "consistent": the abstract supports the claim the sentence attributes to the cited source, as stated. A faithful paraphrase or summary is consistent.
- "contradicts": the abstract conflicts with that claim. Examples: the opposite direction of an effect or finding; a different number; a different study design, method or measure; a different population; a claim stated more strongly, more generally or with more certainty than the abstract reports.
- "silent": the abstract neither supports nor contradicts that claim, so the abstract alone cannot settle it. The claim may still be true in the full paper. Background or methodological uses the abstract does not address are silent.
- "undecidable": only if the sentence or abstract is too garbled or ambiguous for any defensible verdict.

Judge only the claim attributed to THIS source (the sentence may cite other sources too). Use only the abstract; do not use knowledge of the paper beyond it. The reviewers' verdicts are arguments, not votes: majority does not decide, the abstract does.

CITED SOURCE
Title: {title}
Year: {year}
Abstract:
\"\"\"
{abstract}
\"\"\"

CITING SENTENCE (the citation to judge: {cited_as})
"{sentence}"

REVIEWERS
{reviewers}

Return only this JSON object:
{{"verdict": "consistent" | "contradicts" | "silent" | "undecidable",
  "evidence_quote": "<for consistent/contradicts: the shortest span copied exactly, word for word, from the abstract that decides it; otherwise empty>",
  "rationale": "<a detailed explanation, several sentences, of what the sentence claims about this source, what the abstract does and does not say, and why that yields your verdict>",
  "reviewer_assessment": {{"A": "<where reviewer A is right or wrong, and why>", "B": "...", "C": "..."}}}}
"""


def _reviewer_block(order: list[str], seats_v: dict) -> str:
    out = []
    for letter, seat in zip("ABC", order):
        v = seats_v[seat]
        verdict = v.get("guarded") or f"no usable verdict (answered {v.get('raw_verdict')!r}; its quote was not found in the abstract or the answer was malformed)"
        out.append(f"Reviewer {letter}: verdict {verdict}\n  quote: \"{v.get('evidence_quote') or ''}\"\n  reason: {v.get('reason') or ''}")
    return "\n".join(out)


def _decider_items() -> list[dict]:
    """The 76 non-unanimous items and the pre-registered 10% unanimity audit (same seeded draw)."""
    cons = json.loads((STAGE_E / "panel" / "consensus.json").read_text(encoding="utf-8"))
    rows = cons["items"]
    jobs = {jf.stem: json.loads(jf.read_text(encoding="utf-8")) for jf in (STAGE_E / "jobs").glob("*.json")}
    meta = {m["id"]: m for j in jobs.values() for m in j["pairs_meta"]}
    strata = select_for_owner(rows, {cid: meta[cid]["discipline"] for cid in meta})
    out = []
    for r in rows:
        st = strata.get(r["citation_id"], [])
        if "a_no_consensus" in st:
            out.append({**r, "purpose": "adjudicate"})
        elif "b_unanimous_audit" in st:
            out.append({**r, "purpose": "unanimity_audit"})
    return out


def adjudicate(limit: int | None) -> None:
    import random
    out = STAGE_E / "adjudication" / "decider.jsonl"
    out.parent.mkdir(parents=True, exist_ok=True)
    done = set()
    if out.exists():
        for line in out.read_text(encoding="utf-8").splitlines():
            r = json.loads(line)
            if r.get("final"):
                done.add(r["citation_id"])
    jobs = {jf.stem: json.loads(jf.read_text(encoding="utf-8")) for jf in (STAGE_E / "jobs").glob("*.json")}
    cit = {c["id"]: c for j in jobs.values() for c in j["citations"]}
    meta = {m["id"]: m for j in jobs.values() for m in j["pairs_meta"]}
    absd = {a["doi"]: a for j in jobs.values() for a in j["abstracts"]}
    seats = {s: _latest(s) for s in SEATS}
    items = [it for it in _decider_items() if it["citation_id"] not in done]
    if limit:
        items = items[:limit]
    meta_run = {"cli_version": hermetic.cli_version("claude"), "argv": DECIDER_ARGV,
                "prompt_template_sha256": _sha(DECIDER_PROMPT)}
    print(f"[decider] {len(items)} items to decide ({len(done)} done)")
    for k, it in enumerate(items, 1):
        cid = it["citation_id"]
        c, m = cit[cid], meta[cid]
        a = absd[c["doi"]]
        order = list(SEATS)
        random.Random(f"{SEED}|{cid}").shuffle(order)
        sv = {s: seats[s][cid] for s in SEATS}
        prompt = DECIDER_PROMPT.format(title=a.get("title") or "(unknown)", year=a.get("year") or "n.d.",
                                       abstract=a["abstract"], cited_as=m["cited_ref"], sentence=c["sentence"],
                                       reviewers=_reviewer_block(order, sv))
        final = None
        for attempt in (1, 2):
            argv, kw = hermetic.spawn_kwargs(DECIDER_ARGV)
            t0 = time.time()
            cp = subprocess.run(argv, input=prompt, capture_output=True, text=True, encoding="utf-8",
                                errors="replace", timeout=1800, **kw)
            rec = {**meta_run, "at": _now(), "citation_id": cid, "purpose": it["purpose"], "attempt": attempt,
                   "seat_order": dict(zip("ABC", order)), "prompt_sha256": _sha(prompt),
                   "seconds": round(time.time() - t0, 1), "exit": cp.returncode}
            text, problem = "", None
            try:
                d = json.loads(cp.stdout)
                text = d.get("result") or ""
                rec["resolved_models"] = list((d.get("modelUsage") or {}).keys())
                rec["usage"] = d.get("usage")
                if d.get("is_error"):
                    problem = "cli is_error"
            except json.JSONDecodeError:
                problem = "cli output not JSON"
            rec["raw_text"] = text
            parsed = None
            if not problem:
                try:
                    parsed = extract_json_object(text)
                except (ValueError, json.JSONDecodeError):
                    problem = "answer has no JSON object"
            if parsed is not None:
                v = str(parsed.get("verdict") or "").strip().lower()
                if v not in DECIDER_VERDICTS:
                    problem = f"verdict {v!r} outside the scheme"
                elif v in ("consistent", "contradicts") and not scorer.quote_in_abstract(a["abstract"], parsed.get("evidence_quote")):
                    problem = "evidence quote not found in the abstract"
                rec["parsed"] = parsed
            rec["problem"] = problem
            rec["final"] = problem is None or attempt == 2
            rec["decision"] = (str(parsed.get("verdict")).strip().lower() if problem is None
                               else ("undecidable" if rec["final"] else None))
            with out.open("a", encoding="utf-8") as f:
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")
            if problem is None:
                final = rec
                break
            time.sleep(10)
        print(f"[decider] {k}/{len(items)} {it['purpose']} panel={[sv[s]['guarded'] for s in SEATS]} -> "
              f"{(final or rec)['decision']} ({rec['seconds']} s)", flush=True)


def _decisions() -> dict:
    out = {}
    p = STAGE_E / "adjudication" / "decider.jsonl"
    for line in p.read_text(encoding="utf-8").splitlines() if p.exists() else []:
        r = json.loads(line)
        if r.get("final"):
            out[r["citation_id"]] = r
    return out


def build_gold() -> None:
    """gold/gold_mmc_v1.jsonl + gold/DECISIONS.md + gold/summary.json, all in custody."""
    import math
    cons = json.loads((STAGE_E / "panel" / "consensus.json").read_text(encoding="utf-8"))
    dec = _decisions()
    need = {it["citation_id"]: it["purpose"] for it in _decider_items()}
    missing = [c for c in need if c not in dec]
    if missing:
        raise SystemExit(f"{len(missing)} decider items not decided yet; run `adjudicate` first")
    jobs = {jf.stem: json.loads(jf.read_text(encoding="utf-8")) for jf in (STAGE_E / "jobs").glob("*.json")}
    cit = {c["id"]: c for j in jobs.values() for c in j["citations"]}
    meta = {m["id"]: m for j in jobs.values() for m in j["pairs_meta"]}
    absd = {a["doi"]: a for j in jobs.values() for a in j["abstracts"]}
    seats = {s: _latest(s) for s in SEATS}
    gold, audit = [], []
    for r in cons["items"]:
        cid = r["citation_id"]
        d = dec.get(cid)
        if r["unanimous"]:
            label, source = r["consensus"], "frontier_consensus"
            if d:
                audit.append((cid, r["consensus"], d["decision"]))
        else:
            label, source = d["decision"], "adjudicated"
        gold.append({"citation_id": cid, "job": r["job"], "doi": r["doi"], "label": label,
                     "label_source": source, "panel": r["panel"], "free_model": r["free_model"],
                     "discipline": meta[cid]["discipline"], "journal": meta[cid]["journal"],
                     "sentence_sources": meta[cid]["sentence_sources"],
                     "decider": ({"decision": d["decision"], "purpose": d["purpose"],
                                  "rationale": (d.get("parsed") or {}).get("rationale"),
                                  "evidence_quote": (d.get("parsed") or {}).get("evidence_quote"),
                                  "reviewer_assessment": (d.get("parsed") or {}).get("reviewer_assessment"),
                                  "seat_order": d["seat_order"], "resolved_models": d.get("resolved_models"),
                                  "problem": d.get("problem")} if d else None)})
    gdir = STAGE_E / "gold"
    gdir.mkdir(parents=True, exist_ok=True)
    (gdir / "gold_mmc_v1.jsonl").write_text("".join(json.dumps(g, ensure_ascii=False) + "\n" for g in gold),
                                            encoding="utf-8")

    def wilson(k, n):
        if not n:
            return None
        z, p = 1.96, k / n
        den = 1 + z * z / n
        c = (p + z * z / (2 * n)) / den
        h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
        return [round(c - h, 3), round(c + h, 3)]

    agree = sum(1 for _, c, d in audit if c == d)
    fm = {}
    for want in ("contradicts", "consistent"):
        rows = [g for g in gold if g["free_model"]["verdict"] == want and g["label"] != "undecidable"]
        if want == "consistent":
            key = json.loads((STAGE_E / "review" / "stage_e_review_key_r1.json").read_text(encoding="utf-8"))
            sample = {k["citation_id"] for k in key["rows"] if "c_free_consistent_sample" in k["strata"]}
            rows = [g for g in rows if g["citation_id"] in sample]
        k = sum(1 for g in rows if g["label"] == want)
        fm[want] = {"n": len(rows), "agree": k, "rate": round(k / len(rows), 3) if rows else None,
                    "wilson95": wilson(k, len(rows)), "labels": dict(Counter(g["label"] for g in rows))}
    summary = {"built_at": _now(), "items": len(gold),
               "label_source": dict(Counter(g["label_source"] for g in gold)),
               "labels": dict(Counter(g["label"] for g in gold)),
               "labels_by_source": {s: dict(Counter(g["label"] for g in gold if g["label_source"] == s))
                                    for s in ("frontier_consensus", "adjudicated")},
               "adjudicated_vs_panel": dict(Counter(
                   "decider sided with a majority" if sum(v == g["label"] for v in g["panel"].values()) >= 2
                   else "decider sided with one seat" if any(v == g["label"] for v in g["panel"].values())
                   else "decider chose a verdict no seat gave"
                   for g in gold if g["label_source"] == "adjudicated")),
               "seat_agreement_with_label": {s: f"{sum(1 for g in gold if g['panel'][s] == g['label'])}/{len(gold)}"
                                             for s in SEATS},
               "unanimity_audit": {"n": len(audit), "decider_agrees": agree, "wilson95": wilson(agree, len(audit)),
                                   "disagreements": [{"citation_id": c, "consensus": x, "decider": y}
                                                     for c, x, y in audit if x != y]},
               "free_model_agreement_with_frontier_labels": fm}
    (gdir / "summary.json").write_text(json.dumps(summary, indent=1), encoding="utf-8")

    seat_name = {"opus": "Opus 5.5", "sonnet": "Sonnet 5.5", "sol": "GPT-6.1 Sol"}
    md = ["# citeclaim Stage E: decision log for the MetaMisCitations frontier-adjudicated labels", "",
          f"Built {summary['built_at']}. Custody only (article text). Process: ScienceArena",
          "`docs/design/2026-10-citeclaim-stage-e-mmc-pilot.md`, amendments 1 and 1a.", "",
          f"Items: {len(gold)}; label source {summary['label_source']}; labels {summary['labels']}.",
          f"Unanimity audit: decider agrees with the consensus on {agree}/{len(audit)}.", ""]
    for g in sorted(gold, key=lambda g: (g["label_source"] != "adjudicated", g["citation_id"])):
        d = g["decider"]
        if not d:
            continue
        c, m = cit[g["citation_id"]], meta[g["citation_id"]]
        a = absd[c["doi"]]
        md += [f"## {g['citation_id']} ({'ADJUDICATED' if g['label_source'] == 'adjudicated' else 'unanimity audit'})", "",
               f"- Citing paper: {m['journal']}; cited as *{m['cited_ref']}*; source: {a.get('title')} ({a.get('year')}), "
               f"https://doi.org/{c['doi']}",
               f"- Sentence: “{c['sentence']}”",
               f"- Abstract: {a['abstract']}", ""]
        for s in SEATS:
            v = seats[s][g["citation_id"]]
            letter = next(L for L, x in d["seat_order"].items() if x == s)
            md.append(f"- {seat_name[s]} (shown as Reviewer {letter}): **{v['guarded'] or 'no usable verdict'}**"
                      f" — quote “{v.get('evidence_quote') or ''}”; reason: {v.get('reason') or ''}")
        md += ["", f"**Decision: {d['decision']}**" + (f" ({d['problem']})" if d.get("problem") else ""),
               f"- Evidence quote: “{d.get('evidence_quote') or ''}”",
               f"- Rationale: {d.get('rationale') or ''}"]
        for L, txt in (d.get("reviewer_assessment") or {}).items():
            md.append(f"- On Reviewer {L} ({seat_name.get(d['seat_order'].get(L), '?')}): {txt}")
        md += [f"- Label used: **{g['label']}** ({g['label_source']})", ""]
    (gdir / "DECISIONS.md").write_text("\n".join(md), encoding="utf-8")
    print(json.dumps({k: summary[k] for k in summary if k != "built_at"}, indent=1))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["panel", "consensus", "count", "sheet", "adjudicate", "gold"])
    ap.add_argument("--seat", choices=list(SEATS))
    ap.add_argument("--limit", type=int)
    a = ap.parse_args()
    if a.cmd == "panel":
        if not a.seat:
            ap.error("--seat is required")
        panel(a.seat, a.limit)
    elif a.cmd == "consensus":
        consensus()
    elif a.cmd == "sheet":
        sheet()
    elif a.cmd == "adjudicate":
        adjudicate(a.limit)
    elif a.cmd == "gold":
        build_gold()
    else:
        gs = groups()
        print(json.dumps({"calls": len(gs), "items": sum(len(g["citation_ids"]) for g in gs),
                          "group_sizes": dict(Counter(len(g["citation_ids"]) for g in gs))}))


if __name__ == "__main__":
    main()
